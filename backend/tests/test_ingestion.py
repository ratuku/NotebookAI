from io import BytesIO

import pytest
from notebook_ai.pipeline import RagPipeline
from notebook_ai.providers import AIError, AnswerDraft, Claim
from pptx import Presentation
from pptx.util import Inches
from qdrant_client import QdrantClient

from notebook_backend.config import settings
from notebook_backend.controllers import notebooks
from notebook_backend.services import ingestion


class TestProvider:
    model_key = "integration-test"
    fail = False

    def embed(self, texts):
        if self.fail:
            raise AIError("Embedding unavailable")
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        return [1.0, 0.0]

    def generate(self, question, evidence):
        return AnswerDraft(
            insufficient_evidence=False,
            claims=[Claim(text=evidence[0]["text"], citations=[evidence[0]["id"]])],
        )


@pytest.fixture()
def rag(client, monkeypatch):
    pipeline = RagPipeline(QdrantClient(":memory:"), TestProvider())
    monkeypatch.setattr(settings, "ai_mode", "rag")
    monkeypatch.setattr(ingestion, "get_pipeline", lambda: pipeline)
    monkeypatch.setattr(notebooks, "get_pipeline", lambda: pipeline)
    yield pipeline
    pipeline.client.close()


def test_upload_indexes_immediately_and_answers(client, notebook, rag):
    base = f"/api/v1/notebooks/{notebook}"
    response = client.post(
        base + "/sources/upload", files={"file": ("plants.txt", b"Plants use sunlight.")}
    )
    assert response.status_code == 201
    source = response.json()
    assert source["index_status"] == "ready" and source["chunk_count"] == 1
    answer = client.post(
        base + "/questions", json={"question": "sunlight", "source_ids": [source["id"]]}
    )
    assert answer.status_code == 201
    assert answer.json()["mode"] == "rag"
    assert answer.json()["citations"][0]["chunk_id"]


def test_failure_is_saved_and_retry_recovers(client, notebook, rag):
    base = f"/api/v1/notebooks/{notebook}"
    rag.provider.fail = True
    source = client.post(
        base + "/sources/text", json={"title": "Notes", "text": "Plants use sunlight."}
    ).json()
    assert source["index_status"] == "failed"
    assert "Embedding unavailable" in source["index_error"]
    assert (
        client.post(
            base + "/questions", json={"question": "sun", "source_ids": [source["id"]]}
        ).status_code
        == 409
    )
    rag.provider.fail = False
    response = client.post(base + f"/sources/{source['id']}/ingest", json={})
    assert response.json()["index_status"] == "ready"
    assert response.json()["index_error"] is None


def test_rebuild_from_postgres_and_transcription_replaces_old_index(client, notebook, rag):
    base = f"/api/v1/notebooks/{notebook}"
    source = client.post(
        base + "/sources/upload", files={"file": ("notes.txt", b"Old content.")}
    ).json()
    rag.client.delete_collection(rag.collection)
    rebuilt = client.post("/api/v1/ingestion/rebuild", json={"notebook_id": notebook}).json()
    assert rebuilt[0]["index_status"] == "ready"
    assert rag.client.count(rag.collection).count == 1
    changed = client.put(
        base + f"/sources/{source['id']}/text", json={"title": "New", "text": "New content."}
    )
    assert changed.json()["index_status"] == "ready"
    client.post("/api/v1/ingestion/rebuild", json={"notebook_id": notebook})
    records, _ = rag.client.scroll(rag.collection)
    assert [r.payload["text"] for r in records] == ["New content."]
    assert client.delete(base + f"/sources/{source['id']}").status_code == 204
    assert rag.client.count(rag.collection).count == 0


def test_pptx_and_url_have_slide_and_paragraph_citations(client, notebook, rag, monkeypatch):
    base = f"/api/v1/notebooks/{notebook}"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Photosynthesis"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Plants", "Sunlight"
    data = BytesIO()
    presentation.save(data)
    source = client.post(
        base + "/sources/upload", files={"file": ("lecture.pptx", data.getvalue())}
    ).json()
    detail = client.get(base + f"/sources/{source['id']}").json()
    assert detail["segments"][0]["location"] == "Slide 1"
    assert "Plants | Sunlight" in detail["text"]
    monkeypatch.setattr(
        notebooks,
        "fetch_page",
        lambda url: (
            url,
            b"<html><title>Plant guide</title><nav>Ignore this</nav><main><p>Plants use sunlight.</p></main></html>",
            "text/html",
        ),
    )
    web = client.post(base + "/sources/url", json={"url": "https://example.com/plants"}).json()
    assert web["kind"] == "url" and web["index_status"] == "ready"
    assert (
        client.get(base + f"/sources/{web['id']}/file")
        .headers["content-type"]
        .startswith("text/plain")
    )
    assert "Ignore this" not in client.get(base + f"/sources/{web['id']}").json()["text"]


def test_ingestion_endpoint_rejects_other_notebooks(client, notebook, rag):
    other = client.post("/api/v1/notebooks", json={"title": "Other"}).json()["id"]
    source = client.post(
        f"/api/v1/notebooks/{other}/sources/text",
        json={"title": "Other", "text": "Other evidence."},
    ).json()
    assert (
        client.post(
            f"/api/v1/notebooks/{notebook}/sources/{source['id']}/ingest", json={}
        ).status_code
        == 404
    )
