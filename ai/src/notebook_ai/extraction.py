import warnings
from io import BytesIO
from pathlib import Path

from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader


class ExtractionError(ValueError):
    pass


MAX_TEXT = 200_000


def text_segments(text: str) -> list[dict]:
    # Bound each citation; splitting long paragraphs also makes retrieval more useful.
    segments = []
    for paragraph, value in enumerate(text.splitlines(), 1):
        for start in range(0, len(value), 1000):
            chunk = value[start : start + 1000].strip()
            if chunk:
                segments.append({"location": f"Paragraph {paragraph}", "text": chunk})
    return segments


def parse_upload(filename: str, data: bytes) -> dict:
    extension = Path(filename).suffix.lower()
    if extension == ".txt":
        try:
            text = data.decode("utf-8-sig").strip()
        except UnicodeDecodeError:
            raise ExtractionError("Text files must use UTF-8 encoding.")
        if not text or "\x00" in text or len(text) > MAX_TEXT:
            raise ExtractionError("Provide non-empty text, up to 200,000 characters.")
        return {
            "kind": "text",
            "status": "ready",
            "text": text,
            "segments": text_segments(text),
            "media_type": "text/plain",
        }
    if extension == ".pptx":
        return parse_pptx(data)
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("This file is not a valid PDF.")
        try:
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionError("Upload an unencrypted PDF.")
            if len(reader.pages) > 100:
                raise ExtractionError("PDFs may contain at most 100 pages.")
            segments = []
            total = 0
            for number, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or "").replace("\x00", "").strip()
                total += len(text)
                if total > MAX_TEXT:
                    raise ExtractionError("PDF text exceeds 200,000 characters.")
                segments.extend(
                    {"location": f"Page {number}", "text": chunk["text"]}
                    for chunk in text_segments(text)
                )
        except ExtractionError:
            raise
        except Exception:  # noqa: BLE001 — normalize third-party parser errors at this boundary
            raise ExtractionError("Could not read this PDF. Try another file.")
        return {
            "kind": "pdf",
            "status": "ready" if segments else "needs_text",
            "text": "\n\n".join(s["text"] for s in segments),
            "segments": segments,
            "media_type": "application/pdf",
        }
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(data)) as img:
                    if img.format not in {"PNG", "JPEG", "WEBP"}:
                        raise ValueError("Unsupported image")
                    media_type = Image.MIME[img.format]
                    img.verify()
        except (
            UnidentifiedImageError,
            OSError,
            ValueError,
            Image.DecompressionBombWarning,
            Image.DecompressionBombError,
        ):
            raise ExtractionError("This file is not a supported, valid image.")
        return {
            "kind": "image",
            "status": "needs_text",
            "text": "",
            "segments": [],
            "media_type": media_type,
        }
    raise ExtractionError("Supported files: PDF, PPTX, TXT, PNG, JPG, and WEBP.")


def parse_pptx(data: bytes) -> dict:
    from zipfile import ZipFile

    from pptx import Presentation

    try:
        with ZipFile(BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 2000
                or sum(i.file_size for i in archive.infolist()) > 50_000_000
            ):
                raise ExtractionError("The expanded presentation is too large.")
            if "ppt/presentation.xml" not in archive.namelist():
                raise ExtractionError("This file is not a PPTX presentation.")
        presentation = Presentation(BytesIO(data))
        if len(presentation.slides) > 100:
            raise ExtractionError("Presentations may contain at most 100 slides.")

        def shape_text(shapes):
            result = []
            for shape in sorted(shapes, key=lambda item: (item.top, item.left)):
                if shape.has_text_frame:
                    result.extend(p.text for p in shape.text_frame.paragraphs)
                if shape.has_table:
                    result.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                if hasattr(shape, "shapes"):
                    result.extend(shape_text(shape.shapes))
            return result

        segments = []
        total = 0
        for number, slide in enumerate(presentation.slides, 1):
            text = "\n".join(shape_text(slide.shapes)).replace("\x00", "").strip()
            total += len(text)
            if total > MAX_TEXT:
                raise ExtractionError("Presentation text exceeds 200,000 characters.")
            if text:
                segments.append({"location": f"Slide {number}", "text": text})
        return {
            "kind": "pptx",
            "status": "ready" if segments else "needs_text",
            "text": "\n\n".join(s["text"] for s in segments),
            "segments": segments,
            "media_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            "Could not read this PPTX. Upload a valid, unencrypted presentation."
        ) from exc
